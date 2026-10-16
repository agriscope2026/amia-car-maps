"""End-to-end orchestration: inputs → validation → ENVI → map → presentation → outputs/<timestamp>/."""
from __future__ import annotations

import json
import shutil
import zipfile
from datetime import datetime
from pathlib import Path

import pandas as pd

from . import config, index, io_loader, mapping, presentation, recommendations


def merged_settings(settings: dict | None) -> dict:
    s = json.loads(json.dumps(config.DEFAULT_SETTINGS))
    for k, v in (settings or {}).items():
        if v is not None:
            s[k] = v
    return s


def load_inputs(files: dict, settings: dict) -> tuple[dict, io_loader.Gazetteer]:
    """files: {dataset: (path_or_bytes, filename)}"""
    gaz = io_loader.load_gazetteer()
    return io_loader.load_all(files, gaz, settings), gaz


def default_files() -> dict:
    return {k: (p, p.name) for k, p in config.DEFAULT_FILES.items()}


def validation_ok(results: dict) -> bool:
    return all(r.ok for r in results.values())


# ------------------------------------------------------------------ export ---

def match_report(results: dict) -> pd.DataFrame:
    rows = []
    for ds, r in results.items():
        for m in r.report.get("rows", []):
            rows.append({"Dataset": ds, "Row": m.get("row"), "Input name": m.get("raw_name"), "Role": m.get("role"),
                         "Province hint": m.get("province_hint", ""), "Matched GIS area": m.get("area_label") or "",
                         "Area key": m.get("area_key") or "", "Method": m.get("method", ""),
                         "Score": m.get("score", "")})
    return pd.DataFrame(rows)


def methodology_rows(meta: dict, settings: dict, results: dict) -> list[tuple[str, str]]:
    w = meta["weights"]
    rows = [
        ("Title", config.APP_TITLE),
        ("Outlook period", meta["labels"]["period"]),
        ("Issued (as of)", meta["labels"]["issued"]),
        ("Standing crops as of", meta["labels"]["crops_as_of"]),
        *([("Livestock inventory as of", meta["labels"]["livestock_as_of"])]
          if meta.get("exposure", {}).get("livestock") else []),
        ("Analysis level", meta["level"]),
        ("Areas", str(meta["n_areas"])),
        ("Areas with no data", str(meta["n_nodata"])),
        ("Normalization", "Min–max: x_norm = (x − min) / (max − min); 0 when max = min. Higher = more vulnerable."),
        ("Drought forecast", meta["hazard_note"] + f" Normalized with: {meta['hazard_method']}."),
        ("Historical damage", meta["damage_note"]),
        (meta.get("crops_label", "Standing crops"), meta["crops_note"]),
        ("Weights", f"Drought forecast {w['hazard']:.4f}; Historical damage {w['damage']:.4f}; "
                    f"{meta.get('crops_label', 'Standing crops')} {w['crops']:.4f}"),
        ("ENVI", "Σ weight × normalized indicator"),
        ("Classification", meta["classification"]),
        ("Class breaks", " | ".join(f"{b:.4f}" for b in meta["breaks"])),
        ("Missing-data policy", settings.get("missing_policy", "nodata")),
    ]
    cm = results["drought"].report.get("category_mapping")
    if cm:
        rows.append(("Drought category mapping", "; ".join(f"{k} = {v}" for k, v in cm.items())))
    for ds, r in results.items():
        rows.append((f"Input file ({ds})", r.report.get("filename", "")))
    return rows


def export_excel(df: pd.DataFrame, meta: dict, results: dict, settings: dict, path: Path) -> Path:
    table = index.display_table(df)
    with pd.ExcelWriter(path, engine="xlsxwriter") as xw:
        wb = xw.book
        F = config.FONT_FAMILY
        wb.formats[0].set_font_name(F)        # workbook default font
        hdr = wb.add_format({"bold": True, "bg_color": "#0b6237", "font_color": "white", "border": 1,
                             "text_wrap": True, "valign": "vcenter", "font_name": F})
        n3 = wb.add_format({"num_format": "0.000", "font_name": F})
        n2 = wb.add_format({"num_format": "#,##0.00", "font_name": F})
        table.to_excel(xw, sheet_name="ENVI Results", index=False)
        ws = xw.sheets["ENVI Results"]
        for i, c in enumerate(table.columns):
            ws.write(0, i, c, hdr)
            fmt = n3 if "norm" in c or c == "ENVI" else (n2 if "(ha" in c or c.startswith(("crop_", "dmg_")) else None)
            ws.set_column(i, i, 22 if c in ("Area", "Province", "Missing indicators") else 14, fmt)
        ws.set_row(0, 34)
        ws.freeze_panes(1, 3)
        ws.autofilter(0, 0, len(table), len(table.columns) - 1)
        ci = list(table.columns).index("Class")
        for name, color in {**config.CLASS_COLORS, config.NO_DATA_LABEL: config.NO_DATA_COLOR}.items():
            ws.conditional_format(1, ci, len(table), ci, {"type": "cell", "criteria": "==", "value": f'"{name}"',
                                                          "format": wb.add_format({"bg_color": color, "font_name": F})})

        write_computation_sheet(xw.book, df, meta, settings, results, hdr)
        write_monthly_sheet(xw, df, meta, settings, results, hdr)

        s = meta["summary"]
        summ = pd.DataFrame([{"Class": it["class"], "Range": it["range"], "Areas": it["count"]} for it in s["legend"]])
        summ.to_excel(xw, sheet_name="Summary", index=False)
        ws2 = xw.sheets["Summary"]
        for i, c in enumerate(summ.columns):
            ws2.write(0, i, c, hdr)
        ws2.set_column(0, 0, 18)
        ws2.set_column(1, 2, 16)
        r0 = len(summ) + 2
        ws2.write(r0, 0, "Standing crop area at risk (High + Very High), ha", wb.add_format({"bold": True, "font_name": F}))
        ws2.write(r0, 2, s["crop_at_risk_ha"], n2)
        for j, (k, v) in enumerate(s["crop_at_risk_by_crop"].items()):
            ws2.write(r0 + 1 + j, 0, f"  {k}")
            ws2.write(r0 + 1 + j, 2, v, n2)
        r1 = r0 + 2 + len(s["crop_at_risk_by_crop"]) + 1
        top = pd.DataFrame(s["top10"])[["rank", "area", "province", "envi", "class", "crops_ha"]]
        top.columns = ["Rank", "Area", "Province", "ENVI", "Class", "Standing crops (ha)"]
        ws2.write(r1, 0, "Top 10 most vulnerable areas", wb.add_format({"bold": True, "font_name": F}))
        top.to_excel(xw, sheet_name="Summary", index=False, startrow=r1 + 1)
        for i, c in enumerate(top.columns):
            ws2.write(r1 + 1, i, c, hdr)
        ws2.set_column(3, 5, 16)

        meth = pd.DataFrame(methodology_rows(meta, settings, results), columns=["Item", "Value"])
        meth.to_excel(xw, sheet_name="Methodology", index=False)
        ws3 = xw.sheets["Methodology"]
        ws3.set_column(0, 0, 28)
        ws3.set_column(1, 1, 110, wb.add_format({"text_wrap": True, "valign": "top", "font_name": F}))
        for i, c in enumerate(meth.columns):
            ws3.write(0, i, c, hdr)

        mr = match_report(results)
        mr.to_excel(xw, sheet_name="Name Matching", index=False)
        ws4 = xw.sheets["Name Matching"]
        for i, c in enumerate(mr.columns):
            ws4.write(0, i, c, hdr)
        ws4.set_column(0, 8, 18)

        write_recommendations_sheet(xw.book, settings.get("recommendations", ""), hdr)
    return path


def write_computation_sheet(wb, df: pd.DataFrame, meta: dict, settings: dict, results: dict, hdr) -> None:
    """'Computation' sheet: raw indicators → min–max normalization → weighted sum → class, as live Excel formulas
    (with the computed values cached), so the normalization can be checked and weights tried out in Excel."""
    from xlsxwriter.utility import xl_rowcol_to_cell as cell

    F = config.FONT_FAMILY
    ws = wb.add_worksheet("Computation")
    bold = wb.add_format({"bold": True, "font_name": F})
    note = wb.add_format({"italic": True, "font_color": "#555555", "font_name": F})
    n3 = wb.add_format({"num_format": "0.000", "font_name": F})
    n4 = wb.add_format({"num_format": "0.0000", "font_name": F})
    inp = wb.add_format({"num_format": "0.0000", "bg_color": "#fff2cc", "border": 1, "font_name": F})
    nk = config.NO_DATA_LABEL
    rep = results["drought"].report
    fixed = settings.get("hazard_scaling", "minmax") == "fixed" and rep.get("hazard_type") == "categorical"
    invert = not fixed and rep.get("hazard_direction") == "lower_worse"
    zero_missing = settings.get("missing_policy", "nodata") == "zero"
    livestock = "livestock_raw" in df and "livestock_norm" in df

    ws.write(0, 0, "ENVI computation – live Excel formulas", bold)
    ws.write(1, 0, "Change a weight (yellow) or a raw value and the normalized values, ENVI and class recalculate. "
                   "ENVI = w_D × Drought_norm + w_H × Damage_norm + w_C × Crops_norm.", note)

    # parameter block: one row per indicator (weight, min, max, method)
    first, n = 11, len(df)                         # data rows start at Excel row 12
    last = first + n - 1
    raw_col = {"hazard": 3, "damage": 4, "crops": 5}
    labels = {"hazard": "Drought forecast score", "damage": "Historical damage (ha, mean/yr)",
              "crops": "Crop area (ha)" if livestock else "Standing crops (ha)"}
    for i, h in enumerate(["Indicator", "Weight", "Min", "Max", "Normalization"]):
        ws.write(3, i, h, hdr)
    prow = {}
    for j, k in enumerate(index.IND_KEYS):
        r = 4 + j
        prow[k] = r
        rng = f"{cell(first, raw_col[k], True, False)}:{cell(last, raw_col[k], True, False)}"
        vals = pd.to_numeric(df[f"{k}_raw"], errors="coerce")
        lo, hi = vals.min(), vals.max()
        ws.write(r, 0, labels[k])
        ws.write_number(r, 1, float(meta["weights"][k]), inp)
        ws.write_formula(r, 2, f"=MIN({rng})", n4, None if pd.isna(lo) else float(lo))
        if k == "hazard" and fixed:
            ws.write_number(r, 3, float(rep.get("hazard_max") or 0), n4)
            method = "score ÷ maximum possible score (Max column), capped at 0–1"
        else:
            ws.write_formula(r, 3, f"=MAX({rng})", n4, None if pd.isna(hi) else float(hi))
            method = ("min–max inverted: (max − x) / (max − min)" if k == "hazard" and invert
                      else "min–max: (x − min) / (max − min); 0 when max = min")
        ws.write(r, 4, method)
    ws.write(7, 0, "Sum of weights", bold)
    ws.write_formula(7, 1, "=SUM(B5:B7)", n4, float(sum(meta["weights"].values())))

    if livestock:   # exposure = (1 − share) × crop area norm + share × livestock norm (row 10)
        lv = pd.to_numeric(df["livestock_raw"], errors="coerce")
        lrng = f"{cell(first, 11, True, False)}:{cell(last, 11, True, False)}"
        ws.write(9, 0, "Livestock share of exposure")
        ws.write_number(9, 1, float(meta["exposure"]["livestock_share"]), inp)
        ws.write_formula(9, 2, f"=MIN({lrng})", n4, float(lv.min()))
        ws.write_formula(9, 3, f"=MAX({lrng})", n4, float(lv.max()))
        ws.write(9, 4, "Livestock (TLU) min–max; Exposure = (1 − share) × Crop area_norm + share × Livestock_norm")

    br = [float(b) for b in meta["breaks"]]
    ws.write(8, 0, "Class breaks", bold)
    for i, b in enumerate(br):
        ws.write_number(8, 1 + i, b, n4)
    ws.write(8, 1 + len(br), "equal intervals: lower bound included" if meta["classification"] == "equal"
             else "upper bound included", note)

    heads = ["Rank", "Area", "Province", labels["hazard"], labels["damage"], labels["crops"],
             "Drought (norm)", "Damage (norm)", "Exposure (norm)" if livestock else "Crops (norm)", "ENVI", "Class"]
    if livestock:
        heads += ["Livestock (TLU)", "Livestock (norm)"]
    hr = first - 1
    for i, h in enumerate(heads):
        ws.write(hr, i, h, hdr)
    ws.set_row(hr, 34)

    blank = "0" if zero_missing else '""'
    cmp = "<" if meta["classification"] == "equal" else "<="
    names = config.CLASS_NAMES
    for i, (_, row) in enumerate(df.iterrows()):
        r = first + i
        ws.write(r, 1, row["area"])
        ws.write(r, 2, row["province_label"])
        for k, c in raw_col.items():
            v = row[f"{k}_raw"]
            if pd.notna(v):
                ws.write_number(r, c, float(v))
        for k, c in raw_col.items():
            x = cell(r, c)
            mn, mx = cell(prow[k], 2, True, True), cell(prow[k], 3, True, True)
            if k == "hazard" and fixed:
                f = f"IF({mx}=0,0,MIN(1,MAX(0,{x}/{mx})))"
            elif k == "hazard" and invert:
                f = f"IF({mx}={mn},0,({mx}-{x})/({mx}-{mn}))"
            else:
                f = f"IF({mx}={mn},0,({x}-{mn})/({mx}-{mn}))"
            if k == "crops" and livestock:
                lx, lmn, lmx = cell(r, 11), "$C$10", "$D$10"
                lvf = f"IF({lmx}={lmn},0,({lx}-{lmn})/({lmx}-{lmn}))"
                f = f"(1-$B$10)*{f}+$B$10*{lvf}"
                if pd.notna(row["livestock_raw"]):
                    ws.write_number(r, 11, float(row["livestock_raw"]))
                ws.write_formula(r, 12, f'=IF({lx}="",0,{lvf})', n3, float(row["livestock_norm"])
                                 if pd.notna(row["livestock_norm"]) else 0.0)
            v = row[f"{k}_norm"]
            ws.write_formula(r, c + 3, f'=IF({x}="",{blank},{f})', n3, "" if pd.isna(v) else float(v))
        g, h, k3, e = cell(r, 6), cell(r, 7), cell(r, 8), cell(r, 9)
        envi = row["envi"]
        ws.write_formula(r, 9, f'=IF(OR({g}="",{h}="",{k3}=""),"",$B$5*{g}+$B$6*{h}+$B$7*{k3})', n3,
                         "" if pd.isna(envi) else float(envi))
        cls = f'"{names[-1]}"'
        for ci in range(len(names) - 2, -1, -1):          # nested IFs against the inner breaks C9…F9
            cls = f'IF({e}{cmp}{cell(8, 2 + ci, True, True)},"{names[ci]}",{cls})'
        ws.write_formula(r, 10, f'=IF({e}="","{nk}",{cls})', None, row["envi_class"])
        ws.write_formula(r, 0, f'=IF({e}="","",RANK({e},$J${first + 1}:$J${last + 1}))', None,
                         "" if pd.isna(row["rank"]) else int(row["rank"]))

    ws.set_column(0, 0, 24)
    ws.set_column(1, 2, 20)
    ws.set_column(3, 9, 14)
    ws.set_column(10, 10, 12)
    ws.freeze_panes(first, 3)
    ws.autofilter(hr, 0, last, len(heads) - 1)
    for name, color in {**config.CLASS_COLORS, nk: config.NO_DATA_COLOR}.items():
        ws.conditional_format(first, 10, last, 10, {"type": "cell", "criteria": "==", "value": f'"{name}"',
                                                     "format": wb.add_format({"bg_color": color, "font_name": F})})


def write_monthly_sheet(xw, df: pd.DataFrame, meta: dict, settings: dict, results: dict, hdr) -> None:
    """'Monthly ENVI' sheet: per outlook month the drought score, its scaled value (score ÷ 3), the ENVI and class."""
    months = index.monthly_envi(df, meta, settings, results["drought"].report)
    if not months:
        return
    table = df[["area", "province_label"]].rename(columns={"area": "Area", "province_label": "Province"})
    for m, f in months.items():
        table[f"{m} drought score"] = f["hazard_raw"]
        table[f"{m} drought (÷3)"] = f["hazard_norm"].round(3)
        table[f"{m} ENVI"] = f["envi"].round(3)
        table[f"{m} class"] = f["envi_class"]
    table.to_excel(xw, sheet_name="Monthly ENVI", index=False, startrow=2)
    ws = xw.sheets["Monthly ENVI"]
    F = config.FONT_FAMILY
    ws.write(0, 0, "Monthly ENVI = w_D × (month drought score ÷ 3) + w_H × Damage_norm + w_C × Crops_norm "
                   "(damage and crops as in the season ENVI)", xw.book.add_format({"bold": True, "font_name": F}))
    for i, c in enumerate(table.columns):
        ws.write(2, i, c, hdr)
    ws.set_row(2, 34)
    ws.set_column(0, 1, 20)
    ws.set_column(2, len(table.columns) - 1, 11)
    ws.freeze_panes(3, 2)


def write_recommendations_sheet(wb, text: str, hdr=None) -> None:
    """Add a 'Recommendations' sheet (xlsxwriter workbook)."""
    if not (text or "").strip():
        return
    ws = wb.add_worksheet("Recommendations")
    ws.set_column(0, 0, 140, wb.add_format({"text_wrap": True, "valign": "top", "font_name": config.FONT_FAMILY}))
    ws.write(0, 0, "Agricultural recommendations", hdr or wb.add_format({"bold": True, "font_name": config.FONT_FAMILY}))
    for i, line in enumerate(text.splitlines(), start=1):
        ws.write(i, 0, line)


def _replace_recommendations_xlsx(path: Path, text: str) -> None:
    """Rewrite the Recommendations sheet of an existing results workbook (openpyxl)."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = load_workbook(path)
    if "Recommendations" in wb.sheetnames:
        del wb["Recommendations"]
    if text.strip():
        ws = wb.create_sheet("Recommendations")
        ws.column_dimensions["A"].width = 140
        ws["A1"] = "Agricultural recommendations"
        ws["A1"].font = Font(name=config.FONT_FAMILY, bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor="0B6237")
        for i, line in enumerate(text.splitlines(), start=2):
            c = ws.cell(i, 1, line)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.font = Font(name=config.FONT_FAMILY)
    wb.save(path)


def _zip_run(out: Path) -> Path:
    zpath = out / f"ENVI_outputs_{out.name}.zip"
    zpath.unlink(missing_ok=True)
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for p in sorted(out.rglob("*")):
            if p.is_file() and p != zpath:
                z.write(p, p.relative_to(out))
    return zpath


# --------------------------------------------------------------------- run ---

def new_run_dir() -> Path:
    base = config.OUTPUTS / datetime.now().strftime("%Y-%m-%d_%H%M")
    path, n = base, 2
    while path.exists():
        path = base.with_name(f"{base.name}_{n}")
        n += 1
    path.mkdir(parents=True)
    return path


def run(files: dict, settings: dict | None = None, progress=None, out_dir: Path | None = None) -> dict:
    """Full pipeline. ``progress(step, state, message)`` is called for UI updates.

    Returns a manifest with the run folder and produced files.
    """
    progress = progress or (lambda *a: None)
    settings = merged_settings(settings)
    out = out_dir or new_run_dir()

    progress("validate", "running", "Validating inputs")
    results, gaz = load_inputs(files, settings)
    if not validation_ok(results):
        errs = {k: r.report["errors"] for k, r in results.items() if r.report["errors"]}
        progress("validate", "error", json.dumps(errs))
        raise ValueError(f"Input validation failed: {errs}")
    inp = out / "inputs"
    inp.mkdir(exist_ok=True)
    for ds, (src, name) in files.items():
        target = inp / name
        if isinstance(src, (bytes, bytearray)):
            target.write_bytes(src)
        else:
            shutil.copy2(src, target)
    progress("validate", "done", "Inputs validated and saved")

    progress("compute", "running", "Normalizing and computing ENVI")
    df, meta = index.compute_envi(results, gaz, settings)
    if not (settings.get("recommendations") or "").strip():
        settings["recommendations"] = recommendations.generate(df, meta, results, settings)
        settings["recommendations_source"] = "auto"
    xlsx = export_excel(df, meta, results, settings, out / "envi_results.xlsx")
    index.display_table(df).to_csv(out / "envi_results.csv", index=False, encoding="utf-8-sig")
    match_report(results).to_csv(out / "match_report.csv", index=False, encoding="utf-8-sig")
    progress("compute", "done", f"ENVI computed for {meta['n_areas']} areas")

    progress("map", "running", "Rendering map")
    maps = mapping.render(df, meta, out, settings, log=lambda m: progress("map", "running", m))
    areas, prov = maps.pop("areas"), maps.pop("prov")
    (out / "envi_results.geojson").write_text(json.dumps(mapping.to_geojson(areas, prov, 0)["areas"]),
                                              encoding="utf-8")
    progress("map", "done", f"Map rendered ({maps['renderer']})")

    progress("presentation", "running", "Building presentation images")
    pres = presentation.build_all(meta, maps["map_only"], out, settings.get("recommendations", ""))
    progress("presentation", "done", "Presentation images ready")

    progress("package", "running", "Saving settings and ZIP")
    settings_out = {**settings, "renderer_used": maps["renderer"], "run_folder": out.name,
                    "created": datetime.now().isoformat(timespec="seconds"),
                    "input_files": {k: v[1] for k, v in files.items()},
                    "class_breaks": meta["breaks"], "analysis_level": meta["level"],
                    "methodology": dict(methodology_rows(meta, settings, results))}
    (out / "settings.json").write_text(json.dumps(settings_out, indent=2, ensure_ascii=False), encoding="utf-8")
    (out / "meta.json").write_text(json.dumps(meta, indent=1, ensure_ascii=False, default=str), encoding="utf-8")
    zpath = _zip_run(out)
    progress("package", "done", "All outputs saved")

    return {
        "run_id": out.name, "folder": str(out), "renderer": maps["renderer"],
        "files": {
            "excel": xlsx.name, "csv": "envi_results.csv", "map_png": maps["png"].name, "map_pdf": maps["pdf"].name,
            "map_only": maps["map_only"].name, "slide": pres["slide"].name, "poster": pres["poster"].name,
            "zip": zpath.name, "settings": "settings.json", "geojson": "envi_results.geojson",
            "gpkg": maps["gpkg"].name, "match_report": "match_report.csv",
        },
        "summary": meta["summary"],
        "recommendations": settings.get("recommendations", ""),
    }


def rebuild_presentation(run_id: str, text: str) -> dict:
    """Re-draw the presentation of an existing run with edited recommendations (no map re-render)."""
    out = config.OUTPUTS / run_id
    meta = json.loads((out / "meta.json").read_text(encoding="utf-8"))
    pres = presentation.build_all(meta, out / "envi_map_only.png", out, text)
    sp = out / "settings.json"
    st = json.loads(sp.read_text(encoding="utf-8"))
    st["recommendations"] = text
    st["recommendations_source"] = "edited"
    sp.write_text(json.dumps(st, indent=2, ensure_ascii=False), encoding="utf-8")
    if (out / "envi_results.xlsx").exists():
        _replace_recommendations_xlsx(out / "envi_results.xlsx", text)
    _zip_run(out)
    return {k: v.name for k, v in pres.items()}


def list_runs() -> list[dict]:
    if not config.OUTPUTS.exists():
        return []
    runs = []
    for d in sorted(config.OUTPUTS.iterdir(), reverse=True):
        if d.is_dir() and not d.name.startswith("_") and (d / "settings.json").exists():
            runs.append({"run_id": d.name, "files": sorted(p.name for p in d.iterdir() if p.is_file())})
    return runs
