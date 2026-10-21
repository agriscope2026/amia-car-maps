"""Excel template for the banner programs (Rice, Corn, HVCDP, Livestock, …) to submit farm recommendations per weather
condition – tropical cyclone wind signal, heavy rainfall, rainfall forecast class, drought, ENVI class, … – with
translations into Ilocano and another local language.

Sheets:
  * HOW TO FILL – instructions;
  * RECOMMENDATIONS – one recommendation per row; dropdowns for program, weather, condition (depends on the weather
    chosen), timing, stage and language;
  * LISTS – the dropdown values. The conditions use the same class names as the maps of the portal (classes.py,
    pagasa_tc.py), so the rows can later be matched to a map automatically.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from .classes import PRESETS
from .envi import config

PROGRAMS = ["Rice", "Corn", "High-Value Crops (HVCDP)", "Livestock and Poultry", "Organic Agriculture",
            "Fisheries", "All programs (general)"]


def _classes(preset: str) -> list[str]:
    out = []
    for c in PRESETS[preset]["classes"]:
        desc = c.get("desc")
        out.append(f"{desc} ({c['label']} mm)" if desc else c["label"])
    return out


# weather / hazard → (short name for the dropdown list, conditions)
WEATHER: list[tuple[str, str, list[str]]] = [
    ("Tropical cyclone (wind signal)", "TC", ["Tropical cyclone outside PAR / no signal yet (watch)",
                                              "Signal No. 1", "Signal No. 2", "Signal No. 3", "Signal No. 4",
                                              "Signal No. 5"]),
    ("Heavy rainfall", "HEAVY_RAIN", ["Forecast rain 50–100 mm (PAGASA Weather Advisory)",
                                      "Forecast rain 100–200 mm (PAGASA Weather Advisory)",
                                      "Forecast rain > 200 mm (PAGASA Weather Advisory)",
                                      "Yellow rainfall warning", "Orange rainfall warning", "Red rainfall warning"]),
    ("Thunderstorm / strong gusts", "STORM", ["Thunderstorm advisory", "Strong to gale-force gusts"]),
    ("10-day rainfall forecast", "DEKAD", _classes("rainfall_dekad")),
    ("Monthly rainfall forecast", "MONTH", _classes("rainfall_month")),
    ("Seasonal rainfall (% of normal)", "PCT_NORMAL", _classes("pct_normal")),
    ("Drought / dry spell", "DROUGHT", [c["label"] for c in PRESETS["drought"]["classes"] if c["label"] != "Not affected"]),
    ("El Niño vulnerability (ENVI)", "ENVI", list(config.CLASS_NAMES)),
    ("La Niña", "LA_NINA", ["La Niña watch", "La Niña alert / advisory"]),
    ("Low temperature / frost", "COLD", ["Cold temperature (below 10 °C)", "Ground frost"]),
    ("Any weather (general)", "GENERAL", ["Any condition"]),
]
TIMING = ["Before (prepare)", "During", "After (recover)", "Anytime"]
STAGES = ["Any stage", "Land preparation", "Seedbed / nursery", "Planting / transplanting", "Vegetative",
          "Flowering / reproductive", "Maturity / ready to harvest", "Post-harvest / storage",
          "All animals", "Young animals / newborn", "Breeding / pregnant animals", "Fishpond / cages"]
LANGUAGES = ["Kankanaey", "Ibaloi", "Ifugao (Tuwali)", "Kalinga", "Itneg", "Isnag", "Filipino", "Other (write it in REMARKS)"]

COLUMNS = [  # header, width, note
    ("NO.", 6, "Row number (optional)."),
    ("BANNER PROGRAM", 24, "Choose from the list."),
    ("COMMODITY", 22, "Crop or animal, e.g. Palay, Cabbage, Coffee, Swine, Native chicken. Leave blank if it applies to the whole program."),
    ("WEATHER / HAZARD", 30, "Choose from the list."),
    ("CONDITION / LEVEL", 44, "Choose from the list (it shows the levels of the weather chosen in the previous column)."),
    ("WHEN", 16, "Before, during or after the event."),
    ("STAGE", 26, "Crop or animal stage it applies to."),
    ("RECOMMENDATION (ENGLISH)", 60, "One action per row, short and clear, starting with a verb. E.g. 'Harvest palay that is at least 85% ripe.'"),
    ("TRANSLATION – ILOCANO", 60, "The same recommendation in Ilocano."),
    ("OTHER LANGUAGE", 20, "Optional: a second local language."),
    ("TRANSLATION – OTHER LANGUAGE", 60, "The recommendation in the language chosen in the previous column."),
    ("PRIORITY", 10, "1 = most important. Used to choose what fits on the map's recommendation page."),
    ("SOURCE / REFERENCE", 28, "E.g. PhilRice Rice Production Guide, BPI technoguide, BAI."),
    ("SUBMITTED BY", 26, "Office / banner program focal person."),
    ("REMARKS", 30, ""),
]

EXAMPLES = [
    [1, "Rice", "Palay", "Tropical cyclone (wind signal)", "Signal No. 1", "Before (prepare)", "Maturity / ready to harvest",
     "Harvest palay that is at least 85% ripe; start with low-lying fields.", "", "", "", 1,
     "PhilRice Rice Production Guide", "Rice Program", "EXAMPLE – delete this row"],
    [2, "High-Value Crops (HVCDP)", "Highland vegetables", "Heavy rainfall", "Forecast rain 100–200 mm (PAGASA Weather Advisory)",
     "Before (prepare)", "Any stage", "Clear canals and drainage around the beds so water does not stand.", "", "", "", 2,
     "", "HVCDP", "EXAMPLE – delete this row"],
    [3, "Livestock and Poultry", "Swine", "Drought / dry spell", "Dry spell", "During", "All animals",
     "Give clean drinking water at all times and shade the pens.", "", "", "", 1, "", "Livestock Program",
     "EXAMPLE – delete this row"],
]

GREEN = "0B6237"
LIGHT = "EAF3EC"
thin = Side(style="thin", color="C9D3CD")


def build() -> bytes:
    wb = Workbook()

    # --- HOW TO FILL
    how = wb.active
    how.title = "HOW TO FILL"
    lines = [
        ("Farm recommendations per weather condition – DA-RFO-CAR AMIA", "title"),
        ("For the banner programs (Rice, Corn, HVCDP, Livestock, Organic Agriculture, …). The recommendations are used on the "
         "maps and advisories of the CAR agri-climate portal (10-day and seasonal rainfall, drought, ENVI and the Special Farm "
         "Weather Outlook and Advisory for tropical cyclones).", ""),
        ("", ""),
        ("How to fill it in", "head"),
        ("1. Use the sheet RECOMMENDATIONS. One row = one recommendation (one action).", ""),
        ("2. Choose the BANNER PROGRAM, the WEATHER / HAZARD, then its CONDITION / LEVEL from the dropdowns. Choose the weather "
         "first – the condition list depends on it.", ""),
        ("3. Choose WHEN it applies (before, during or after the event) and the STAGE of the crop or animal.", ""),
        ("4. Write the RECOMMENDATION in English: short, clear and doable by farmers, starting with a verb. Avoid long "
         "explanations – one action per row. Add more rows for more actions.", ""),
        ("5. Write the TRANSLATION in Ilocano. If you can, add a second local language (Kankanaey, Ibaloi, Ifugao, Kalinga, …) "
         "in OTHER LANGUAGE and its translation.", ""),
        ("6. PRIORITY: 1 for the most important actions. Maps have limited space; the highest priorities are shown first.", ""),
        ("7. Delete the example rows, then send the file back to the AMIA Program.", ""),
        ("", ""),
        ("Weather conditions (as on the maps)", "head"),
    ]
    for w, _, conds in WEATHER:
        lines.append((f"• {w}: " + "; ".join(conds), ""))
    lines += [("", ""), ("Tips", "head"),
              ("• Tropical cyclones: the higher the signal, the more urgent and the shorter the action. Signal No. 1 = "
               "prepare (36 hours before strong winds); Signal No. 3–5 = only quick, safe actions – people first.", ""),
              ("• Do not change the headers or the dropdown values: the portal reads them.", "")]
    how.column_dimensions["A"].width = 130
    for i, (text, kind) in enumerate(lines, start=1):
        c = how.cell(row=i, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "title":
            c.font = Font(bold=True, size=15, color=GREEN)
        elif kind == "head":
            c.font = Font(bold=True, size=12, color=GREEN)

    # --- LISTS (dropdown values; one named range per weather for the dependent CONDITION list)
    lists = wb.create_sheet("LISTS")
    blocks = [("PROGRAMS", PROGRAMS), ("WEATHER", [w for w, _, _ in WEATHER]), ("TIMING", TIMING), ("STAGES", STAGES),
              ("LANGUAGES", LANGUAGES)] + [(f"C_{code}", conds) for _, code, conds in WEATHER]
    for col, (name, values) in enumerate(blocks, start=1):
        letter = get_column_letter(col)
        lists.cell(row=1, column=col, value=name).font = Font(bold=True)
        for r, v in enumerate(values, start=2):
            lists.cell(row=r, column=col, value=v)
        lists.column_dimensions[letter].width = max(14, min(60, max(len(v) for v in values) + 2))
        wb.defined_names[name] = DefinedName(name, attr_text=f"LISTS!${letter}$2:${letter}${len(values) + 1}")
    # weather → name of its condition list
    col = len(blocks) + 2
    lists.cell(row=1, column=col, value="WEATHER").font = Font(bold=True)
    lists.cell(row=1, column=col + 1, value="CONDITION LIST").font = Font(bold=True)
    for r, (w, code, _) in enumerate(WEATHER, start=2):
        lists.cell(row=r, column=col, value=w)
        lists.cell(row=r, column=col + 1, value=f"C_{code}")
    a, b = get_column_letter(col), get_column_letter(col + 1)
    wb.defined_names["WEATHER_MAP"] = DefinedName("WEATHER_MAP", attr_text=f"LISTS!${a}$2:${b}${len(WEATHER) + 1}")
    lists.column_dimensions[a].width = 32
    lists.column_dimensions[b].width = 16
    lists.sheet_state = "hidden"

    # --- RECOMMENDATIONS
    ws = wb.create_sheet("RECOMMENDATIONS", 1)
    head_fill = PatternFill("solid", fgColor=GREEN)
    for i, (h, width, note) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = width
        n = ws.cell(row=2, column=i, value=note)
        n.font = Font(italic=True, size=9, color="5F6368")
        n.fill = PatternFill("solid", fgColor=LIGHT)
        n.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 64
    ws.freeze_panes = "C3"
    first, last = 3, 1002
    for r, row in enumerate(EXAMPLES, start=first):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v if v != "" else None)
            cell.font = Font(color="7A7A7A", italic=True)
    for r in range(first, last + 1):
        for c in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=thin)

    def dv(formula: str, col: str, prompt: str, strict: bool = True):
        v = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False,
                           showErrorMessage=strict, errorStyle="stop" if strict else "warning",
                           error="Choose a value from the list.", promptTitle=COLUMNS[ord(col) - 65][0], prompt=prompt)
        v.showInputMessage = True
        ws.add_data_validation(v)
        v.add(f"{col}{first}:{col}{last}")

    dv("PROGRAMS", "B", "Choose the banner program.")
    dv("WEATHER", "D", "Choose the weather or hazard first.")
    # the condition list of the weather in column D of the same row
    dv(f"INDIRECT(VLOOKUP($D{first},WEATHER_MAP,2,FALSE))", "E", "Choose the level of the weather in column D.")
    dv("TIMING", "F", "Before, during or after the event.")
    dv("STAGES", "G", "Crop or animal stage.", strict=False)
    dv("LANGUAGES", "J", "Second local language (optional).", strict=False)
    pr = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                        error="Priority: a whole number from 1 (most important) to 5.", showErrorMessage=True)
    ws.add_data_validation(pr)
    pr.add(f"L{first}:L{last}")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"
    wb.active = 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
